from io import BytesIO
from xml.sax.saxutils import escape

from django.db import models, transaction
from django.db.models import Q
from django.http import FileResponse, HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from rest_framework import serializers, status, viewsets
from rest_framework.decorators import action
from rest_framework.response import Response

from .models import (
	Entrada, EntradaDetalle, Inventario, Lote, MovimientoInventario, OrdenCompra,
	OrdenCompraDetalle, Salida, SalidaDetalle,
)
from .serializers import (
	EntradaDetalleSerializer, EntradaSerializer, InventarioSerializer, LoteSerializer,
	MovimientoInventarioSerializer, OrdenCompraDetalleSerializer, OrdenCompraSerializer,
	SalidaDetalleSerializer, SalidaSerializer,
)


class OrdenCompraViewSet(viewsets.ModelViewSet):
	queryset = OrdenCompra.objects.select_related('proveedor', 'creado_por').prefetch_related('detalles__producto').all()
	serializer_class = OrdenCompraSerializer

	def get_queryset(self):
		queryset = super().get_queryset()
		for field in ('proveedor', 'estado'):
			value = self.request.query_params.get(field)
			if value:
				queryset = queryset.filter(**{f'{field}_id' if field == 'proveedor' else field: value})
		return queryset

	def perform_create(self, serializer):
		serializer.save(creado_por=self.request.user if self.request.user.is_authenticated else None)

	@action(detail=True, methods=['post'])
	def aprobar(self, request, pk=None):
		orden = self.get_object()
		if orden.estado != OrdenCompra.Estado.PENDIENTE:
			return Response({'detail': 'Solo se pueden aprobar órdenes pendientes.'}, status=status.HTTP_400_BAD_REQUEST)
		if not orden.detalles.exists():
			return Response({'detail': 'La orden debe tener al menos un producto.'}, status=status.HTTP_400_BAD_REQUEST)
		orden.estado = OrdenCompra.Estado.APROBADA
		orden.save(update_fields=['estado', 'actualizado_en'])
		return Response(self.get_serializer(orden).data)

	@action(detail=True, methods=['post'])
	def cancelar(self, request, pk=None):
		orden = self.get_object()
		if orden.estado in (OrdenCompra.Estado.COMPLETADA, OrdenCompra.Estado.CANCELADA):
			return Response({'detail': 'La orden no se puede cancelar en su estado actual.'}, status=400)
		orden.estado = OrdenCompra.Estado.CANCELADA
		orden.save(update_fields=['estado', 'actualizado_en'])
		return Response(self.get_serializer(orden).data)


class OrdenCompraDetalleViewSet(viewsets.ModelViewSet):
	queryset = OrdenCompraDetalle.objects.select_related('orden_compra', 'producto').all()
	serializer_class = OrdenCompraDetalleSerializer

	def perform_create(self, serializer):
		orden = serializer.validated_data['orden_compra']
		if orden.estado != OrdenCompra.Estado.PENDIENTE:
			raise serializers.ValidationError('Solo se pueden modificar órdenes pendientes.')
		with transaction.atomic():
			serializer.save()
			orden.recalcular_totales()

	def perform_update(self, serializer):
		if serializer.instance.orden_compra.estado != OrdenCompra.Estado.PENDIENTE:
			raise serializers.ValidationError('Solo se pueden modificar órdenes pendientes.')
		with transaction.atomic():
			detalle = serializer.save()
			detalle.orden_compra.recalcular_totales()

	def perform_destroy(self, instance):
		if instance.orden_compra.estado != OrdenCompra.Estado.PENDIENTE:
			raise serializers.ValidationError('Solo se pueden modificar órdenes pendientes.')
		orden = instance.orden_compra
		with transaction.atomic():
			instance.delete()
			orden.recalcular_totales()


class EntradaViewSet(viewsets.ModelViewSet):
	queryset = Entrada.objects.select_related('proveedor', 'orden_compra').prefetch_related('detalles__producto').all()
	serializer_class = EntradaSerializer

	@action(detail=True, methods=['post'])
	def confirmar(self, request, pk=None):
		with transaction.atomic():
			entrada = Entrada.objects.select_for_update().prefetch_related('detalles').get(pk=pk)
			if entrada.estado != Entrada.Estado.PENDIENTE:
				return Response({'detail': 'Solo se pueden confirmar entradas pendientes.'}, status=400)
			orden = None
			if entrada.orden_compra_id:
				orden = OrdenCompra.objects.select_for_update().get(pk=entrada.orden_compra_id)
				if orden.estado not in (OrdenCompra.Estado.APROBADA, OrdenCompra.Estado.PENDIENTE):
					return Response({'detail': 'La orden no está disponible para recepción.'}, status=400)
			for detalle in entrada.detalles.all():
				if orden:
					try:
						orden_detalle = orden.detalles.get(producto_id=detalle.producto_id)
					except OrdenCompraDetalle.DoesNotExist:
						return Response({'detail': f'El producto {detalle.producto_id} no pertenece a la orden.'}, status=400)
					if orden_detalle.cantidad_recibida + detalle.cantidad > orden_detalle.cantidad:
						return Response({'detail': f'La cantidad recibida supera la orden para el producto {detalle.producto_id}.'}, status=400)
					orden_detalle.cantidad_recibida += detalle.cantidad
					orden_detalle.save(update_fields=['cantidad_recibida'])
				inventario, _ = Inventario.objects.select_for_update().get_or_create(producto_id=detalle.producto_id)
				inventario.existencia += detalle.cantidad
				inventario.save(update_fields=['existencia', 'actualizado_en'])
				MovimientoInventario.objects.create(
					producto_id=detalle.producto_id,
					tipo=MovimientoInventario.Tipo.ENTRADA,
					cantidad=detalle.cantidad,
					referencia_tipo='ENTRADA',
					referencia_id=entrada.id,
					lote=detalle.lote,
					usuario=request.user if request.user.is_authenticated else None,
				)
				if detalle.lote:
					lote, _ = Lote.objects.select_for_update().get_or_create(
						producto_id=detalle.producto_id, codigo=detalle.lote,
					)
					lote.cantidad_disponible += detalle.cantidad
					lote.save(update_fields=['cantidad_disponible'])
			entrada.estado = Entrada.Estado.CONFIRMADA
			entrada.save(update_fields=['estado'])
			if orden:
				orden.estado = OrdenCompra.Estado.COMPLETADA if not orden.detalles.filter(cantidad_recibida__lt=models.F('cantidad')).exists() else OrdenCompra.Estado.APROBADA
				orden.save(update_fields=['estado', 'actualizado_en'])
		return Response(self.get_serializer(entrada).data)


class EntradaDetalleViewSet(viewsets.ModelViewSet):
	queryset = EntradaDetalle.objects.select_related('entrada', 'producto').all()
	serializer_class = EntradaDetalleSerializer

	def perform_create(self, serializer):
		entrada = serializer.validated_data['entrada']
		if entrada.estado != Entrada.Estado.PENDIENTE:
			raise serializers.ValidationError('Solo se pueden modificar entradas pendientes.')
		with transaction.atomic():
			serializer.save()
			entrada.recalcular_totales()


class SalidaViewSet(viewsets.ModelViewSet):
	queryset = Salida.objects.select_related('cliente', 'responsable').prefetch_related('detalles__producto').all()
	serializer_class = SalidaSerializer

	def get_queryset(self):
		queryset = super().get_queryset()
		for field in ('cliente', 'tipo', 'estado'):
			value = self.request.query_params.get(field)
			if value:
				queryset = queryset.filter(**{f'{field}_id' if field == 'cliente' else field: value})
		return queryset

	def perform_create(self, serializer):
		serializer.save()

	@action(detail=True, methods=['post'])
	def completar(self, request, pk=None):
		with transaction.atomic():
			salida = Salida.objects.select_for_update().prefetch_related('detalles').get(pk=pk)
			if salida.estado != Salida.Estado.PENDIENTE:
				return Response({'detail': 'Solo se pueden completar salidas pendientes.'}, status=400)
			if not salida.detalles.exists():
				return Response({'detail': 'La salida debe tener al menos un producto.'}, status=400)
			for detalle in salida.detalles.all():
				inventario, _ = Inventario.objects.select_for_update().get_or_create(producto_id=detalle.producto_id)
				if inventario.existencia < detalle.cantidad:
					return Response({'detail': f'Inventario insuficiente para el producto {detalle.producto_id}.'}, status=400)
				inventario.existencia -= detalle.cantidad
				inventario.save(update_fields=['existencia', 'actualizado_en'])
				MovimientoInventario.objects.create(
					producto_id=detalle.producto_id,
					tipo=MovimientoInventario.Tipo.SALIDA,
					cantidad=detalle.cantidad,
					referencia_tipo='SALIDA',
					referencia_id=salida.id,
					usuario=request.user if request.user.is_authenticated else None,
				)
			salida.estado = Salida.Estado.COMPLETADA
			salida.save(update_fields=['estado'])
		return Response(self.get_serializer(salida).data)


class SalidaDetalleViewSet(viewsets.ModelViewSet):
	queryset = SalidaDetalle.objects.select_related('salida', 'producto').all()
	serializer_class = SalidaDetalleSerializer

	def perform_create(self, serializer):
		salida = serializer.validated_data['salida']
		if salida.estado != Salida.Estado.PENDIENTE:
			raise serializers.ValidationError('Solo se pueden modificar salidas pendientes.')
		with transaction.atomic():
			serializer.save()
			salida.recalcular_totales()

	def perform_update(self, serializer):
		if serializer.instance.salida.estado != Salida.Estado.PENDIENTE:
			raise serializers.ValidationError('Solo se pueden modificar salidas pendientes.')
		with transaction.atomic():
			detalle = serializer.save()
			detalle.salida.recalcular_totales()

	def perform_destroy(self, instance):
		if instance.salida.estado != Salida.Estado.PENDIENTE:
			raise serializers.ValidationError('Solo se pueden modificar salidas pendientes.')
		salida = instance.salida
		with transaction.atomic():
			instance.delete()
			salida.recalcular_totales()


class InventarioViewSet(viewsets.ReadOnlyModelViewSet):
	queryset = Inventario.objects.select_related('producto', 'proveedor', 'producto__tipo_producto').all()
	serializer_class = InventarioSerializer

	def get_queryset(self):
		queryset = super().get_queryset()
		proveedor = self.request.query_params.get('proveedor')
		tipo = self.request.query_params.get('tipoProducto') or self.request.query_params.get('tipo')
		descripcion = self.request.query_params.get('descripcion')
		if proveedor and proveedor.lower() != 'todos':
			provider_filter = Q(proveedor__nombre__iexact=proveedor)
			if proveedor.isdigit():
				provider_filter |= Q(proveedor_id=int(proveedor))
			queryset = queryset.filter(provider_filter)
		if tipo and tipo.lower() != 'todos':
			type_filter = Q(producto__tipo_producto__nombre__iexact=tipo)
			if tipo.isdigit():
				type_filter |= Q(producto__tipo_producto_id=int(tipo))
			queryset = queryset.filter(type_filter).distinct()
		if descripcion:
			queryset = queryset.filter(producto__descripcion__icontains=descripcion)
		return queryset.order_by('producto__descripcion')

	@action(detail=False, methods=['get'], url_path='exportar-pdf')
	def exportar_pdf(self, request):
		return _exportar_pdf(
			'INVENTARIO', 'inventario.pdf', _inventario_headers(),
			_inventario_rows(self.get_queryset()),
		)

	@action(detail=False, methods=['get'], url_path='exportar-excel')
	def exportar_excel(self, request):
		return _exportar_excel(
			'Inventario', 'inventario.xlsx', _inventario_headers(),
			_inventario_rows(self.get_queryset()),
		)


class MovimientoInventarioViewSet(viewsets.ReadOnlyModelViewSet):
	queryset = MovimientoInventario.objects.select_related('producto', 'usuario').all()
	serializer_class = MovimientoInventarioSerializer

	def get_queryset(self):
		queryset = super().get_queryset()
		producto = self.request.query_params.get('producto')
		tipo = self.request.query_params.get('tipo')
		if producto:
			queryset = queryset.filter(producto_id=producto)
		if tipo:
			queryset = queryset.filter(tipo=tipo)
		return queryset


class LoteViewSet(viewsets.ModelViewSet):
	queryset = Lote.objects.select_related('producto').all()
	serializer_class = LoteSerializer


def _inventario_rows(queryset):
	return [[
		item.codigo or item.producto.codigo or item.producto_id,
		item.descripcion or item.producto.descripcion,
		item.producto.tipo_producto.nombre,
		item.proveedor.nombre if item.proveedor else '',
		item.saldo_proveedor,
		item.remision,
		item.devolucion,
		item.reposicion,
		item.en_alquiler,
		item.en_bodega,
		item.estado,
	] for item in queryset]


def _inventario_headers():
	return [
		'CÓDIGO', 'DESCRIPCIÓN', 'TIPO', 'PROVEEDOR', 'SALDO PROVEEDOR',
		'REMISIÓN', 'DEVOLUCIÓN', 'REPOSICIÓN', 'EN ALQUILER', 'EN BODEGA', 'ESTADO',
	]


def _exportar_pdf(title, filename, headers, rows):
	buffer = BytesIO()
	styles = getSampleStyleSheet()
	doc = SimpleDocTemplate(
		buffer, pagesize=landscape(A4), rightMargin=10 * mm, leftMargin=10 * mm,
		topMargin=10 * mm, bottomMargin=10 * mm,
	)
	data = [headers] + rows
	table = Table(
		[[Paragraph(escape(str(value)), styles['BodyText']) for value in row] for row in data],
		colWidths=[22 * mm] + [35 * mm] * (len(headers) - 1), repeatRows=1,
	)
	table.setStyle(TableStyle([
		('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#333333')),
		('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
		('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
		('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
		('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
		('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
	]))
	doc.build([Paragraph(title, styles['Title']), Spacer(1, 8 * mm), table])
	buffer.seek(0)
	return FileResponse(buffer, as_attachment=True, filename=filename, content_type='application/pdf')


def _exportar_excel(sheet_title, filename, headers, rows):
	workbook = Workbook()
	worksheet = workbook.active
	worksheet.title = sheet_title
	worksheet.append(headers)
	for cell in worksheet[1]:
		cell.fill = PatternFill('solid', fgColor='333333')
		cell.font = Font(color='FFFFFF', bold=True)
		cell.alignment = Alignment(horizontal='center')
	for row in rows:
		worksheet.append(row)
	for column in worksheet.columns:
		width = max(len(str(cell.value or '')) for cell in column) + 2
		worksheet.column_dimensions[column[0].column_letter].width = min(width, 60)
	output = BytesIO()
	workbook.save(output)
	output.seek(0)
	return HttpResponse(
		output.getvalue(),
		content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
		headers={'Content-Disposition': f'attachment; filename="{filename}"'},
	)
